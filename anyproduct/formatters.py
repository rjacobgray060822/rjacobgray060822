"""Output format helpers for generator results."""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Callable, Dict, Iterable, Mapping

try:  # pragma: no cover - optional dependency
    from docx import Document  # type: ignore
except ImportError:  # pragma: no cover - docx optional
    Document = None  # type: ignore

from .generators.base import GeneratorResult


@dataclass(frozen=True)
class FormattedOutput:
    """Represents a rendered file."""

    filename: str
    content: bytes
    mime_type: str


@dataclass(frozen=True)
class FormatSpec:
    """Metadata describing an output format."""

    slug: str
    label: str
    extension: str
    mime_type: str
    renderer: Callable[[GeneratorResult], bytes]


def _slugify(text: str) -> str:
    text = text.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    text = re.sub(r"-+", "-", text)
    return text.strip("-") or "output"


def _render_markdown(result: GeneratorResult) -> bytes:
    lines = [f"# {result.title}\n"]
    if result.summary:
        lines.append(result.summary + "\n")
    for section in result.sections:
        lines.append(f"## {section.title}\n")
        lines.append(section.content + "\n")
    for table in result.tables:
        headers = " | ".join(table.columns)
        separator = " | ".join(["---"] * len(table.columns))
        lines.append(f"## {table.name}\n")
        lines.append(f"{headers}\n")
        lines.append(f"{separator}\n")
        for row in table.rows:
            row_values = [str(row.get(column, "")) for column in table.columns]
            lines.append(" | ".join(row_values) + "\n")
    return "\n".join(lines).encode("utf-8")


def _render_json(result: GeneratorResult) -> bytes:
    return json.dumps(result.to_dict(), indent=2).encode("utf-8")


def _render_csv(result: GeneratorResult) -> bytes:
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["section", "content"])
    if result.summary:
        writer.writerow(["Summary", result.summary])
    for section in result.sections:
        writer.writerow([section.title, section.content])
    for table in result.tables:
        writer.writerow([])
        writer.writerow([table.name])
        writer.writerow(list(table.columns))
        for row in table.rows:
            writer.writerow([row.get(column, "") for column in table.columns])
    return buffer.getvalue().encode("utf-8")


def _render_docx(result: GeneratorResult) -> bytes:
    if Document is None:  # pragma: no cover - optional dependency guard
        raise RuntimeError(
            "python-docx is required for DOCX output. Install it via 'pip install python-docx'."
        )
    document = Document()
    document.add_heading(result.title, level=0)
    if result.summary:
        document.add_paragraph(result.summary)
    for section in result.sections:
        document.add_heading(section.title, level=1)
        document.add_paragraph(section.content)
    for table in result.tables:
        document.add_heading(table.name, level=1)
        doc_table = document.add_table(rows=1, cols=len(table.columns))
        header_cells = doc_table.rows[0].cells
        for idx, column in enumerate(table.columns):
            header_cells[idx].text = str(column)
        for row in table.rows:
            row_cells = doc_table.add_row().cells
            for idx, column in enumerate(table.columns):
                row_cells[idx].text = str(row.get(column, ""))
    byte_stream = BytesIO()
    document.save(byte_stream)
    return byte_stream.getvalue()


_SUPPORTED_FORMATS: Mapping[str, FormatSpec] = {
    "markdown": FormatSpec(
        slug="markdown",
        label="Markdown",
        extension="md",
        mime_type="text/markdown",
        renderer=_render_markdown,
    ),
    "json": FormatSpec(
        slug="json",
        label="JSON",
        extension="json",
        mime_type="application/json",
        renderer=_render_json,
    ),
    "csv": FormatSpec(
        slug="csv",
        label="CSV",
        extension="csv",
        mime_type="text/csv",
        renderer=_render_csv,
    ),
    "docx": FormatSpec(
        slug="docx",
        label="DOCX",
        extension="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        renderer=_render_docx,
    ),
}


def supported_formats() -> Mapping[str, FormatSpec]:
    return _SUPPORTED_FORMATS


def render_outputs(
    result: GeneratorResult, formats: Iterable[str], *, basename: str | None = None
) -> Dict[str, FormattedOutput]:
    if basename is None:
        basename = _slugify(result.title)
    outputs: Dict[str, FormattedOutput] = {}
    for slug in formats:
        spec = _SUPPORTED_FORMATS.get(slug)
        if spec is None:
            raise KeyError(f"Unsupported format '{slug}'")
        content = spec.renderer(result)
        filename = f"{basename}.{spec.extension}"
        outputs[slug] = FormattedOutput(
            filename=filename,
            content=content,
            mime_type=spec.mime_type,
        )
    return outputs


def write_outputs(
    result: GeneratorResult, formats: Iterable[str], output_dir: Path
) -> Dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    rendered = render_outputs(result, formats)
    written: Dict[str, Path] = {}
    for slug, output in rendered.items():
        target = output_dir / output.filename
        target.write_bytes(output.content)
        written[slug] = target
    return written
